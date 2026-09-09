"""GET /export/compliance.csv — a downloadable audit package, the one thing
Vanta/Drata-style tools are most consistently asked for: something an auditor
or exec can take away instead of only viewing live in-app."""

from __future__ import annotations

import csv
import io


def rows_for(client, org_id):
    resp = client.get("/export/compliance.csv", headers={"authorization": f"org:{org_id}"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]
    return list(csv.reader(io.StringIO(resp.text)))


def test_export_has_one_row_per_control_with_a_header(client, bootstrap, upload):
    org_id, _ = bootstrap(client)
    upload(client, org_id)

    rows = rows_for(client, org_id)
    assert rows[0] == ["framework", "clause", "title", "verdict", "auditor_verdict", "locked",
                       "owner_emails", "open_gaps"]
    assert len(rows) > 1  # header plus at least one control


def test_export_is_org_isolated(client, bootstrap, upload):
    org_a, _ = bootstrap(client)
    org_b, _ = bootstrap(client)
    upload(client, org_a)

    a_rows = rows_for(client, org_a)
    b_rows = rows_for(client, org_b)
    assert len(a_rows) > 1
    assert len(b_rows) == 1  # header only, no controls


# --------------------------------------------------------------------- xlsx


def xlsx_rows(client, org_id, path="/export/compliance.xlsx", **params):
    from openpyxl import load_workbook

    resp = client.get(path, headers={"authorization": f"org:{org_id}"}, params=params)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_compliance_xlsx_opens_and_matches_the_csv_row_count(client, bootstrap, upload):
    org_id, _ = bootstrap(client)
    upload(client, org_id)

    csv_rows = rows_for(client, org_id)
    xl_rows = xlsx_rows(client, org_id)
    assert xl_rows[0] == csv_rows[0]
    assert len(xl_rows) == len(csv_rows)


def test_compliance_xlsx_verdict_filter_trims_rows(client, bootstrap, upload):
    org_id, _ = bootstrap(client)
    upload(client, org_id)  # no model in tests -> every control links FAIL

    all_rows = xlsx_rows(client, org_id)
    passing = xlsx_rows(client, org_id, verdict="PASS")
    failing = xlsx_rows(client, org_id, verdict="FAIL")
    assert len(passing) == 1  # header only — nothing passes without a model
    assert len(failing) == len(all_rows)


def test_compliance_xlsx_scope_matches_csv_scope(client, bootstrap, upload):
    """Scope regression check: xlsx reuses control_details()/list_gaps() —
    exactly the CSV export's authorization path, nothing new — so an auditor
    identity's xlsx export must be the same size as their CSV export."""
    org_id, engagement_id = bootstrap(client)
    upload(client, org_id)
    headers = {"authorization": f"auditor:{engagement_id}"}

    csv_resp = client.get("/export/compliance.csv", headers=headers)
    xlsx_resp = client.get("/export/compliance.xlsx", headers=headers)
    assert csv_resp.status_code == 200 and xlsx_resp.status_code == 200

    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xlsx_resp.content)).active
    xl_rows = [[c.value for c in row] for row in ws.iter_rows()]
    csv_rows_ = list(csv.reader(io.StringIO(csv_resp.text)))
    assert len(xl_rows) == len(csv_rows_)


def test_gaps_xlsx_open_filter(client, bootstrap, upload):
    org_id, _ = bootstrap(client)
    upload(client, org_id)

    open_rows = xlsx_rows(client, org_id, path="/export/gaps.xlsx", status="OPEN")
    assert open_rows[0] == ["framework", "clause", "attribute", "detail", "actual_value",
                            "required_value", "required_action", "status"]
    assert len(open_rows) > 1
    assert all(row[7] == "OPEN" for row in open_rows[1:])


def test_tasks_xlsx_has_header_even_with_no_tasks(client, bootstrap):
    org_id, _ = bootstrap(client)
    rows = xlsx_rows(client, org_id, path="/export/tasks.xlsx")
    assert rows[0] == ["title", "framework", "clause", "required_action", "status", "priority",
                       "owner_email", "due_at"]
    assert len(rows) == 1
